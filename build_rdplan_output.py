# -*- coding: utf-8 -*-
"""Convert LDA 研发规划与实施计划.md -> HTML (BP 同风格) + DOCX 三件套。"""
import os, re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"D:\agent_LDA\LDA_研发规划与实施计划_2026-09-09.md"
OUT_HTML = r"D:\agent_LDA\LDA_研发规划与实施计划_2026-09-09.html"
OUT_DOCX = r"D:\agent_LDA\LDA_研发规划与实施计划_2026-09-09.docx"

CJK_BODY = "宋体"
CJK_HEAD = "黑体"

# ---------------- markdown parse into blocks ----------------
def parse_blocks(lines):
    blocks = []  # (type, payload)
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()
        if s == "":
            i += 1; continue
        # horizontal rule
        if re.fullmatch(r"-{3,}|={3,}|\*{3,}", s):
            blocks.append(("hr", None)); i += 1; continue
        # blockquote (consecutive >)
        if s.startswith(">"):
            ql = []
            while i < n and lines[i].strip().startswith(">"):
                ql.append(re.sub(r"^\s*>\s?", "", lines[i].strip()))
                i += 1
            blocks.append(("quote", "  ".join(ql))); continue
        # table
        if s.startswith("|") and i+1 < n and is_sep(lines[i+1]):
            rows = [split_cells(line)]
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append(split_cells(lines[i])); i += 1
            blocks.append(("table", rows)); continue
        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            blocks.append(("h", (len(m.group(1)), m.group(2).strip()))); i += 1; continue
        # list
        if re.match(r"^\s*([-*]|\d+\.)\s+", line):
            items = []
            while i < n and re.match(r"^\s*([-*]|\d+\.)\s+", lines[i]):
                lm = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", lines[i])
                indent = len(lm.group(1)); ordered = bool(re.match(r"\d+\.", lm.group(2)))
                items.append((indent, ordered, lm.group(3))); i += 1
            blocks.append(("list", items)); continue
        # paragraph
        blocks.append(("p", s)); i += 1
    return blocks

def is_sep(line):
    s = line.strip().strip("|").replace(" ", "")
    return bool(re.fullmatch(r":?-+:?(\|:?-+:?)*", s)) and "-" in s

def split_cells(line):
    s = line.strip()
    if s.startswith("|"): s = s[1:]
    if s.endswith("|"): s = s[:-1]
    return [c.strip() for c in s.split("|")]

# ---------------- inline parsing ----------------
INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\[.+?\]\(.+?\))")

def html_escape(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

def inline_html(text):
    out = []
    pos = 0
    for m in INLINE.finditer(text):
        if pos < m.start():
            out.append(html_escape(text[pos:m.start()]))
        tok = m.group(0)
        if tok.startswith("**"):
            out.append("<b>%s</b>" % html_escape(tok[2:-2]))
        elif tok.startswith("`"):
            out.append("<code>%s</code>" % html_escape(tok[1:-1]))
        elif tok.startswith("["):
            inner = re.match(r"\[(.+?)\]\((.+?)\)", tok)
            if inner:
                out.append('<a href="%s">%s</a>' % (html_escape(inner.group(2)), html_escape(inner.group(1))))
            else:
                out.append(html_escape(tok))
        pos = m.end()
    if pos < len(text):
        out.append(html_escape(text[pos:]))
    return "".join(out)

def add_runs(paragraph, text, base_size=10.5, base_color=None, italic=False, color=None):
    base_color = color if color is not None else base_color
    pos = 0
    for m in INLINE.finditer(text):
        if pos < m.start():
            r = paragraph.add_run(text[pos:m.start()]); _style(r, base_size, base_color, italic=italic)
        tok = m.group(0)
        if tok.startswith("**"):
            r = paragraph.add_run(tok[2:-2]); _style(r, base_size, base_color, bold=True, italic=italic)
        elif tok.startswith("`"):
            r = paragraph.add_run(tok[1:-1]); _style(r, base_size, base_color, latin="Consolas", italic=italic)
        elif tok.startswith("["):
            inner = re.match(r"\[(.+?)\]\((.+?)\)", tok)
            if inner:
                r = paragraph.add_run(inner.group(1)); _style(r, base_size, base_color, color=RGBColor(0x2E,0x75,0xB6), italic=italic)
            else:
                r = paragraph.add_run(tok); _style(r, base_size, base_color, italic=italic)
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:]); _style(r, base_size, base_color, italic=italic)

def _style(run, size, color=None, bold=False, italic=False, latin="Calibri", cjk=CJK_BODY):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), cjk)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color

# ---------------- HTML emit ----------------
HTML_CSS = """  @page { size: A4; margin: 14mm 13mm; }
  * { box-sizing: border-box; }
  body { font-family: "Microsoft YaHei","PingFang SC","Source Han Sans SC",sans-serif;
         font-size: 10pt; line-height: 1.55; color:#111; background:#fff; margin:0; padding:0; }
  .page { max-width: 184mm; margin: 0 auto; }
  h1 { font-size:19pt; margin:0 0 3mm; letter-spacing:.5px; border-bottom:2.2pt solid #111; padding-bottom:2.5mm; }
  .sub { font-size:8.5pt; color:#555; margin:0 0 5mm; line-height:1.6; }
  h2 { font-size:13pt; margin:6mm 0 2.5mm; padding:1.6mm 3mm; background:#111; color:#fff; }
  h3 { font-size:10.8pt; margin:4mm 0 1.5mm; color:#111; border-left:3pt solid #111; padding-left:2.5mm; }
  h4 { font-size:10.2pt; margin:3.5mm 0 1mm; color:#444; padding-left:2.5mm; }
  p { margin:0 0 2mm; text-align:justify; }
  table { width:100%; border-collapse:collapse; margin:2mm 0 3.5mm; font-size:8.8pt; page-break-inside:avoid; }
  th { background:#ececec; border:.6pt solid #999; padding:1.6mm 2mm; text-align:left; font-weight:700; }
  td { border:.6pt solid #bbb; padding:1.5mm 2mm; vertical-align:top; }
  tr:nth-child(even) td { background:#fafafa; }
  .say { background:#f4f6f8; border-left:2.5pt solid #666; padding:2mm 3mm; margin:1.5mm 0 3mm; font-size:9.4pt;
         page-break-inside:avoid; }
  .say p { margin:0 0 1.5mm; }
  .say p:last-child { margin-bottom:0; }
  ul, ol { margin:1mm 0 2.5mm; padding-left:5mm; }
  li { margin-bottom:1mm; }
  .foot { margin-top:5mm; padding-top:2mm; border-top:.6pt solid #999; font-size:8pt; color:#666; }
  @media print { h2, th, .say { -webkit-print-color-adjust:exact; print-color-adjust:exact; } }"""

def emit_html(blocks):
    parts = []
    parts.append("<!DOCTYPE html>")
    parts.append('<html lang="zh-CN">')
    parts.append("<head>")
    parts.append('<meta charset="UTF-8">')
    parts.append("<title>LDA 研发规划与实施计划 · 2026-09-09</title>")
    parts.append("<style>%s</style>" % HTML_CSS)
    parts.append("</head>")
    parts.append('<body><div class="page">')
    parts.append('<h1>LDA 研发规划与实施计划</h1>')
    parts.append('<p class="sub">版本：v0.9.63 · 2026-09-09 · <b>执行视角</b>（姊妹文档：BP 修订版《LDA 商业计划书与研发融资规划》） · 数字基线以 README 顶行权威账本为准。排除项（杜先生明确）：一人公司（OPC）模式仅讨论、不进本规划。</p>')
    first_h1 = True
    for b in blocks:
        t = b[0]
        if t == "h":
            lvl, txt = b[1]
            if lvl == 1:
                if first_h1:  # skip markdown title (already hard-coded above)
                    first_h1 = False
                    continue
                parts.append('<h1>%s</h1>' % inline_html(txt))
            elif lvl == 2:
                parts.append('<h2>%s</h2>' % inline_html(txt))
            elif lvl == 3:
                parts.append('<h3>%s</h3>' % inline_html(txt))
            else:
                parts.append('<h4>%s</h4>' % inline_html(txt))
        elif t == "p":
            parts.append('<p>%s</p>' % inline_html(b[1]))
        elif t == "quote":
            parts.append('<div class="say"><p>%s</p></div>' % inline_html(b[1]))
        elif t == "hr":
            parts.append('<hr style="border:none;border-top:.6pt solid #999;margin:3mm 0;">')
        elif t == "table":
            rows = b[1]
            cells = ['<tr>%s</tr>' % "".join("<th>%s</th>" % inline_html(c) for c in rows[0])]
            for r in rows[1:]:
                cells.append('<tr>%s</tr>' % "".join("<td>%s</td>" % inline_html(c) for c in r))
            parts.append("<table>%s</table>" % "".join(cells))
        elif t == "list":
            is_ol = b[1][0][1]
            tag = "ol" if is_ol else "ul"
            lis = []
            for indent, ordered, content in b[1]:
                prefix = ""
                if content.startswith("[ ] "):
                    prefix = "☑ " if False else "☐ "
                    content = content[4:]
                elif content.startswith("[x] ") or content.startswith("[X] "):
                    prefix = "☑ "
                    content = content[4:]
                lis.append("<li>%s%s</li>" % (prefix, inline_html(content)))
            parts.append("<%s>%s</%s>" % (tag, "".join(lis), tag))
    parts.append('<p class="foot">本文档由 LDA 规划 MD 经 build_rdplan_output.py 自动转制，与 BP 修订版（2026-09-09）配套。关键技术结论均来自真实运行的 LDA 自研代码与确定性验证脚本；外部事实以 README 顶行权威账本为准。</p>')
    parts.append("</div></body></html>")
    return "\n".join(parts)

# ---------------- DOCX emit ----------------
def set_hr(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "999999")
    pbdr.append(bottom); pPr.append(pbdr)

def add_table(doc, rows):
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=1, cols=ncol)
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = rows[0]
    for i in range(ncol):
        cell = t.rows[0].cells[i]; cell.text = ""
        p = cell.paragraphs[0]; txt = hdr[i] if i < len(hdr) else ""
        add_runs(p, txt, base_size=9)
        for r in p.runs: _style(r, 9, cjk=CJK_HEAD, bold=True)
    for row in rows[1:]:
        cells = t.add_row().cells
        for i in range(ncol):
            cell = cells[i]; cell.text = ""
            p = cell.paragraphs[0]; txt = row[i] if i < len(row) else ""
            add_runs(p, txt, base_size=9)
            for r in p.runs: _style(r, 9)
    try:
        width = Inches(6.5)
        for row in t.rows:
            for cell in row.cells:
                cell.width = width / ncol
    except Exception:
        pass
    doc.add_paragraph()

def emit_docx(blocks, path):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"; style.font.size = Pt(10.5)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), CJK_BODY)

    # cover / title
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("LDA 研发规划与实施计划"); _style(r, 22, cjk=CJK_HEAD, bold=True)
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("执行视角 · 与 BP 修订版配套 · 2026-09-09 · v0.9.63"); _style(r2, 10, cjk=CJK_BODY, color=RGBColor(0x80,0x80,0x80))
    doc.add_paragraph()

    for b in blocks:
        t = b[0]
        if t == "h":
            lvl, txt = b[1]
            if lvl == 1:
                pp = doc.add_paragraph(); pp.style = doc.styles["Heading 1"]; add_runs(pp, txt, base_size=16)
            elif lvl == 2:
                pp = doc.add_paragraph(); pp.style = doc.styles["Heading 2"]; add_runs(pp, txt, base_size=14)
            elif lvl == 3:
                pp = doc.add_paragraph(); pp.style = doc.styles["Heading 3"]; add_runs(pp, txt, base_size=12)
            else:
                pp = doc.add_paragraph(); add_runs(pp, txt, base_size=11, bold=True)
        elif t == "p":
            pp = doc.add_paragraph(); add_runs(pp, b[1]); pp.paragraph_format.space_after = Pt(4)
        elif t == "quote":
            pp = doc.add_paragraph(); pp.paragraph_format.left_indent = Inches(0.4)
            add_runs(pp, b[1], base_size=9.5, italic=True, color=RGBColor(0x55,0x55,0x55))
        elif t == "hr":
            set_hr(doc.add_paragraph())
        elif t == "table":
            add_table(doc, b[1])
        elif t == "list":
            for indent, ordered, content in b[1]:
                cb = content
                if cb.startswith("[ ] "): cb = "☐ " + cb[4:]
                elif cb.startswith("[x] ") or cb.startswith("[X] "): cb = "☑ " + cb[4:]
                pp = doc.add_paragraph(style="List Number" if ordered else "List Bullet")
                try: pp.level = indent // 2
                except Exception: pass
                pp.paragraph_format.left_indent = Inches(0.3 + 0.3 * (indent // 2))
                add_runs(pp, cb, base_size=10)
    doc.save(path)

# ---------------- main ----------------
def main():
    with open(SRC, "r", encoding="utf-8") as f:
        lines = f.read().split("\n")
    blocks = parse_blocks(lines)
    # HTML
    html = emit_html(blocks)
    with open(OUT_HTML, "w", encoding="utf-8") as f:
        f.write(html)
    print("WROTE html bytes=", len(html))
    # DOCX
    emit_docx(blocks, OUT_DOCX)
    print("WROTE docx")

if __name__ == "__main__":
    main()
