# -*- coding: utf-8 -*-
"""
从话术手册 Markdown 真源生成 A4 打印版 HTML 与 Word。

用法（CI python）：
  python scripts/build_talk_kit.py

约定：
  - 真源固定为仓库根 LDA_对外话术手册_客户侧与Foundry侧_2026-09-09.md
  - 支持的 Markdown 子集：# / ## / ###、> 引述块、| 表格、- 列表、1. 列表、**粗体**
  - 章节分页规则：h2 序号（从 0 起）在 PAGEBREAK_AT 集合中的，前面插入分页符
"""
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "LDA_对外话术手册_客户侧与Foundry侧_2026-09-09.md"
OUT_HTML = ROOT / "LDA_对外话术手册_打印版.html"
OUT_DOCX = ROOT / "LDA_对外话术手册_客户侧与Foundry侧_2026-09-09.docx"

# 在这些 h2 前强制分页（序号从 0 起）
PAGEBREAK_AT = {2, 3, 4, 5}


# ---------------------------------------------------------------- 解析 Markdown
def normalize_quotes(text: str) -> str:
    """英文直引号成对转换为中文直角引号（逐行，保证配对）。"""
    out = []
    for line in text.split("\n"):
        if line.count('"') >= 2 and line.count('"') % 2 == 0:
            buf, opened = [], True
            for ch in line:
                if ch == '"':
                    buf.append("「" if opened else "」")
                    opened = not opened
                else:
                    buf.append(ch)
            out.append("".join(buf))
        else:
            out.append(line)
    return "\n".join(out)


def parse(text: str):
    """返回 blocks: (kind, payload)。kind ∈ h1/h2/h3/p/q/table/ul/ol"""
    lines = normalize_quotes(text).split("\n")
    blocks, i = [], 0
    para, qparas, qcur = [], [], []
    h2_idx = -1

    def flush_para():
        nonlocal para
        if para:
            blocks.append(("p", " ".join(para)))
            para = []

    def flush_quote():
        nonlocal qparas, qcur
        if qcur:
            qparas.append(" ".join(qcur))
            qcur = []
        if qparas:
            blocks.append(("q", qparas))
            qparas = []

    while i < len(lines):
        raw = lines[i].rstrip()
        s = raw.strip()

        # 表格
        if s.startswith("|") and s.count("|") >= 2:
            flush_para()
            flush_quote()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                rows.append(row)
                i += 1
            rows = [r for r in rows
                    if not all(re.fullmatch(r":?-{2,}:?", c or "") for c in r)]
            blocks.append(("table", rows))
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            flush_para()
            flush_quote()
            lvl, title = len(m.group(1)), m.group(2).strip()
            if lvl == 1:
                blocks.append(("h1", title))
            elif lvl == 2:
                h2_idx += 1
                blocks.append(("h2", (title, h2_idx)))
            elif lvl == 3:
                blocks.append(("h3", title))
            else:
                blocks.append(("h4", title))
            i += 1
            continue

        # 引述（话术块）
        if s.startswith(">"):
            flush_para()
            content = s.lstrip(">").strip()
            if content:
                qcur.append(content)
            else:
                if qcur:
                    qparas.append(" ".join(qcur))
                    qcur = []
            i += 1
            continue

        # 无序列表
        if re.match(r"^[-*]\s+", s):
            flush_para()
            flush_quote()
            items = []
            while i < len(lines) and re.match(r"^[-*]\s+", lines[i].strip()):
                items.append(re.sub(r"^[-*]\s+", "", lines[i].strip()))
                i += 1
            blocks.append(("ul", items))
            continue

        # 有序列表
        if re.match(r"^\d+\.\s+", s):
            flush_para()
            flush_quote()
            items = []
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i].strip()):
                items.append(re.sub(r"^\d+\.\s+", "", lines[i].strip()))
                i += 1
            blocks.append(("ol", items))
            continue

        # 空行 / 分隔线
        if s == "" or re.fullmatch(r"-{3,}|\*{3,}", s):
            flush_para()
            flush_quote()
            i += 1
            continue

        para.append(s)
        i += 1

    flush_para()
    flush_quote()
    return blocks


def plain(s: str) -> str:
    """去掉 Markdown 标记，得到纯文本（用于 docx / HTML 前的清洗）。"""
    s = s.replace("\\*", "\x00")          # 转义星号先占位
    s = s.replace("**", "")               # 粗体标记
    s = re.sub(r"(?<!\*)\*(?!\*)", "", s)  # 残余单星号（斜体标记）
    return s.replace("\x00", "*")


def split_bold(s: str):
    """把 'a **b** c' 切成 [(text, is_bold), ...]"""
    out, pos = [], 0
    for m in re.finditer(r"\*\*(.+?)\*\*", s):
        if m.start() > pos:
            out.append((s[pos:m.start()], False))
        out.append((m.group(1), True))
        pos = m.end()
    if pos < len(s):
        out.append((s[pos:], False))
    return [(t, b) for t, b in out if t]


# ---------------------------------------------------------------- 生成 HTML
def esc(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def html_inline(s: str) -> str:
    buf = []
    for txt, bold in split_bold(s):
        t = esc(plain(txt))
        buf.append(f"<b>{t}</b>" if bold else t)
    return "".join(buf)


CSS = """
  @page { size: A4; margin: 14mm 13mm; }
  * { box-sizing: border-box; }
  body { font-family: "Microsoft YaHei","PingFang SC","Source Han Sans SC",sans-serif;
         font-size: 10pt; line-height: 1.55; color:#111; background:#fff; margin:0; padding:0; }
  .page { max-width: 184mm; margin: 0 auto; }
  h1 { font-size:19pt; margin:0 0 3mm; letter-spacing:.5px; border-bottom:2.2pt solid #111; padding-bottom:2.5mm; }
  .sub { font-size:8.5pt; color:#555; margin:0 0 5mm; line-height:1.6; }
  h2 { font-size:13pt; margin:6mm 0 2.5mm; padding:1.6mm 3mm; background:#111; color:#fff; }
  h2.pagebreak { page-break-before: always; }
  h3 { font-size:10.8pt; margin:4mm 0 1.5mm; color:#111; border-left:3pt solid #111; padding-left:2.5mm; }
  h4 { font-size:10.2pt; margin:3.5mm 0 1mm; color:#444; padding-left:2.5mm; }
  p { margin:0 0 2mm; text-align:justify; }
  table { width:100%; border-collapse:collapse; margin:2mm 0 3.5mm; font-size:8.8pt; page-break-inside:avoid; }
  th { background:#ececec; border:.6pt solid #999; padding:1.6mm 2mm; text-align:left; font-weight:700; }
  td { border:.6pt solid #bbb; padding:1.5mm 2mm; vertical-align:top; }
  tr:nth-child(even) td { background:#fafafa; }
  .say { background:#f4f6f8; border-left:2.5pt solid #666; padding:2mm 3mm; margin:1.5mm 0 3mm; font-size:9.4pt;
         page-break-inside:avoid; }
  .say p { margin:0 0 1.5mm; }
  .say p:last-child { margin-bottom:0; }
  .key { font-weight:700; }
  ul, ol { margin:1mm 0 2.5mm; padding-left:5mm; }
  li { margin-bottom:1mm; }
  .foot { margin-top:5mm; padding-top:2mm; border-top:.6pt solid #999; font-size:8pt; color:#666; }
  .noprint { text-align:center; margin:0 auto 5mm; padding:3mm; background:#fffbe6;
             border:1pt dashed #c90; font-size:9.5pt; }
  @media print {
    .noprint { display:none !important; }
    h2, th, .say { -webkit-print-color-adjust:exact; print-color-adjust:exact; }
  }
"""


def build_html(blocks) -> str:
    # h1 之后紧跟的引述块作为副标题（.sub），不重复渲染成话术块
    sub_lines = []
    if len(blocks) > 1 and blocks[0][0] == "h1" and blocks[1][0] == "q":
        sub_lines = blocks[1][1]
        blocks = [blocks[0]] + list(blocks[2:])

    out = ['<!DOCTYPE html>', '<html lang="zh-CN">', '<head>', '<meta charset="UTF-8">',
           '<title>LDA 对外话术手册 · 打印版</title>', f'<style>{CSS}</style>', '</head>',
           '<body>', '<div class="page">']
    out.append('<div class="noprint">打印提示：按 <b>Ctrl + P</b> → 目标选「另存为 PDF」或直接打印 → '
               '纸张 A4、边距「默认」、<b>勾选「背景图形」</b>。打印时本提示不会出现在纸上。</div>')

    for kind, payload in blocks:
        if kind == "h1":
            out.append(f'<h1>{html_inline(payload)}</h1>')
            if sub_lines:
                out.append('<p class="sub">' +
                           '<br>'.join(html_inline(x) for x in sub_lines) +
                           '</p>')
        elif kind == "h2":
            title, idx = payload
            cls = " class=\"pagebreak\"" if idx in PAGEBREAK_AT else ""
            out.append(f'<h2{cls}>{html_inline(title)}</h2>')
        elif kind == "h3":
            out.append(f'<h3>{html_inline(payload)}</h3>')
        elif kind == "h4":
            out.append(f'<h4>{html_inline(payload)}</h4>')
        elif kind == "p":
            out.append(f'<p>{html_inline(payload)}</p>')
        elif kind == "q":
            out.append('<div class="say">')
            for para in payload:
                out.append(f'<p>{html_inline(para)}</p>')
            out.append('</div>')
        elif kind == "table":
            out.append('<table>')
            for ri, row in enumerate(payload):
                tag = "th" if ri == 0 else "td"
                out.append('<tr>' + ''.join(f'<{tag}>{html_inline(c)}</{tag}>' for c in row) + '</tr>')
            out.append('</table>')
        elif kind == "ul":
            out.append('<ul>' + ''.join(f'<li>{html_inline(x)}</li>' for x in payload) + '</ul>')
        elif kind == "ol":
            out.append('<ol>' + ''.join(f'<li>{html_inline(x)}</li>' for x in payload) + '</ol>')

    out.append('<p class="foot">本册为对外认知装备，不含任何内部资产信息'
               '（服务器地址、令牌、内部路径均未收录）。</p>')
    out.append('</div></body></html>')
    return "\n".join(out)


# ---------------------------------------------------------------- 生成 DOCX
def build_docx(blocks, path: Path):
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.6)
    sec.top_margin = sec.bottom_margin = Cm(1.6)

    st = doc.styles["Normal"]
    st.font.name = "Microsoft YaHei"
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    def shade(el, fill):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), fill)
        el.append(shd)

    def left_border(p, size=18, color="666666"):
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), str(size))
        left.set(qn("w:space"), "4")
        left.set(qn("w:color"), color)
        pbdr.append(left)
        pPr.append(pbdr)

    def add_runs(p, s, size=None, bold_all=False, color=None):
        for txt, b in split_bold(s):
            r = p.add_run(plain(txt))
            r.bold = b or bold_all
            if size:
                r.font.size = Pt(size)
            r.font.name = "Microsoft YaHei"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            if color:
                r.font.color.rgb = color
        return p

    # h1 之后紧跟的引述块已在 HTML 侧作副标题；docx 侧同样跳过，改在这里输出
    if len(blocks) > 1 and blocks[0][0] == "h1" and blocks[1][0] == "q":
        sub_lines = blocks[1][1]
        blocks = [blocks[0]] + list(blocks[2:])
    else:
        sub_lines = []

    for kind, payload in blocks:
        if kind == "h1":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, payload, size=18, bold_all=True)
            for line in sub_lines:
                sp = doc.add_paragraph()
                sp.paragraph_format.space_after = Pt(1)
                add_runs(sp, line, size=8.5, color=RGBColor(0x55, 0x55, 0x55))
        elif kind == "h2":
            title, idx = payload
            if idx in PAGEBREAK_AT:
                doc.add_page_break()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(6)
            shade(p._p, "111111")
            add_runs(p, title, size=12.5, bold_all=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        elif kind == "h3":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            left_border(p, size=24, color="111111")
            p.paragraph_format.left_indent = Cm(0.25)
            add_runs(p, payload, size=10.8, bold_all=True)
        elif kind == "h4":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(7)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Cm(0.25)
            add_runs(p, payload, size=10.2, bold_all=True,
                     color=RGBColor(0x44, 0x44, 0x44))
        elif kind == "p":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, payload)
        elif kind == "q":
            for para in payload:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.left_indent = Cm(0.2)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                shade(p._p, "F4F6F8")
                left_border(p, size=18, color="888888")
                add_runs(p, para, size=9.4)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif kind == "table":
            rows, cols = len(payload), max(len(r) for r in payload)
            t = doc.add_table(rows=rows, cols=cols)
            t.style = "Table Grid"
            for ri, row in enumerate(payload):
                for ci in range(cols):
                    cell = t.cell(ri, ci)
                    txt = row[ci] if ci < len(row) else ""
                    cell.text = ""
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(1)
                    add_runs(p, txt, size=8.8, bold_all=(ri == 0))
                    if ri == 0:
                        shade(cell._tc.get_or_add_tcPr(), "ECECEC")
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif kind in ("ul", "ol"):
            for j, x in enumerate(payload, 1):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.left_indent = Cm(0.7)
                prefix = f"{j}. " if kind == "ol" else "• "
                p.add_run(prefix).font.name = "Microsoft YaHei"
                add_runs(p, x)

    doc.save(str(path))


# ---------------------------------------------------------------- main
def main():
    import sys

    args = sys.argv[1:]
    if args:
        # 传入任意 .md 真源，输出同名 .html / .docx
        src = Path(args[0]).resolve()
        out_html, out_docx = src.with_suffix(".html"), src.with_suffix(".docx")
    else:
        src, out_html, out_docx = SRC, OUT_HTML, OUT_DOCX

    blocks = parse(src.read_text(encoding="utf-8"))
    out_html.write_text(build_html(blocks), encoding="utf-8")
    build_docx(blocks, out_docx)
    n_tab = sum(1 for k, _ in blocks if k == "table")
    print(f"parsed blocks={len(blocks)} tables={n_tab}")
    print(f"  -> {out_html.name}")
    print(f"  -> {out_docx.name}")


if __name__ == "__main__":
    main()
